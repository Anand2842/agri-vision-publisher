#!/usr/bin/env python3
"""Build the Word letterhead template editors use to typeset an article before exporting it
to PDF and uploading it in Admin → Articles.

Usage: python3 scripts/make-letterhead-template.py
Writes public/templates/article-letterhead-template.docx (served at /templates/...).
"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

GREEN = RGBColor(0x2E, 0x51, 0x34)
GREY = RGBColor(0x55, 0x55, 0x55)
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOGO = os.path.join(ROOT, "src", "assets", "logo.png")
OUT = os.path.join(ROOT, "public", "templates", "article-letterhead-template.docx")


def bottom_border(paragraph, color="2E5134", size=12):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": str(size), "w:space": "4", "w:color": color}.items():
        bottom.set(qn(k), v)
    borders.append(bottom)
    p_pr.append(borders)


def page_field(run):
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def styled(paragraph, text, size=None, bold=False, italic=False, color=None):
    run = paragraph.add_run(text)
    run.bold, run.italic = bold, italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return run


doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15
for name, size in (("Title", 18), ("Heading 1", 13), ("Heading 2", 11.5)):
    st = doc.styles[name]
    st.font.name = "Times New Roman"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = GREEN
# Title: magazine-green rule instead of Word's default blue one
for bdr in doc.styles["Title"].element.xpath(".//w:pBdr"):
    bdr.getparent().remove(bdr)

sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.top_margin = sec.bottom_margin = Cm(2.2)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.header_distance = sec.footer_distance = Cm(1.0)

# Letterhead (repeats on every page)
header = sec.header
table = header.add_table(rows=1, cols=2, width=Cm(16.6))
table.autofit = False
logo_cell, text_cell = table.rows[0].cells
for col, w in zip(table.columns, (Cm(1.8), Cm(14.8))):
    col.width = w
    for cell in col.cells:
        cell.width = w
logo_cell.paragraphs[0].add_run().add_picture(LOGO, width=Cm(1.5))
p = text_cell.paragraphs[0]
styled(p, "The Agriculture Popular Article Magazine", size=14, bold=True, color=GREEN)
p = text_cell.add_paragraph()
styled(p, "Vol. [X], Issue [Y]  ·  [Month Year]  ·  ISSN: Applied for  ·  agriculturemagazine.in", size=8.5, color=GREY)
p.paragraph_format.space_after = Pt(0)
rule = header.add_paragraph()
bottom_border(rule)
header.paragraphs[0].paragraph_format.space_after = Pt(0)

# Footer: citation + page number
footer = sec.footer
p = footer.paragraphs[0]
styled(
    p,
    "Cite as: [Authors] ([Year]). [Article title]. The Agriculture Popular Article Magazine, "
    "Vol. [X], No. [Y], pp. [a–b].",
    size=8, italic=True, color=GREY,
)
p = footer.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
styled(p, "Page ", size=8, color=GREY)
page_field(styled(p, "", size=8, color=GREY))

# Body placeholders
bottom_border(doc.add_paragraph("[Article Title in Title Case]", style="Title"), size=6)
p = doc.add_paragraph()
styled(p, "[Author One]¹, [Author Two]², [Author Three]¹", bold=True)
p = doc.add_paragraph()
styled(p, "¹[Department, Institute, City, State]   ²[Department, Institute, City, State]", size=9.5, italic=True)
p = doc.add_paragraph()
styled(p, "Corresponding author: [email address]", size=9.5, color=GREY)
bottom_border(p, color="BBBBBB", size=4)

p = doc.add_paragraph()
styled(p, "Abstract: ", bold=True)
styled(p, "[150–250 words summarising the article.]")
p = doc.add_paragraph()
styled(p, "Keywords: ", bold=True)
styled(p, "[keyword one, keyword two, keyword three]")

for heading, body in (
    ("Introduction", "[Body text. Paste the article here and use Heading 1 / Heading 2 for section titles.]"),
    ("[Section Heading]", "[Body text. Figures: insert the picture, then a caption below it: “Fig. 1. Caption text.”]"),
    ("Conclusion", "[Body text.]"),
    ("References", "[Author, A. (Year). Title. Journal, Volume(Issue), pages.]"),
):
    doc.add_paragraph(heading, style="Heading 1")
    doc.add_paragraph(body)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
styled(
    p,
    "Editor's checklist (delete before export): fill in Vol/Issue/Month in the header and the "
    "citation in the footer (double-click the header/footer to edit) → File → Save As → PDF → "
    "upload it on the article in Admin → Articles.",
    size=8.5, italic=True, color=GREY,
)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print(OUT, os.path.getsize(OUT), "bytes")
